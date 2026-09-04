"""
tests/test_ksp_parser.py — тесты core/ksp_parser.py.

tests/fixtures/ksp_sample_*.docx — три файла с намеренно разной вёрсткой
(одна таблица на весь документ / пять отдельных таблиц / опечатки и
разный регистр в заголовках колонок) — ровно та ситуация, о которой
предупреждает PLAN_STAGE1.md, Б3.2. ksp_sample_1_single_table.doc — тот
же документ, пересохранённый LibreOffice в бинарный .doc, для проверки
Б3.1 на настоящем .doc, а не только на .docx.
"""

import json
from pathlib import Path

import pytest
from docx import Document

from core.db import execute, init_db, query
from core.ksp_parser import (
    KSPConversionError,
    KSPParseError,
    _match_cell_role,
    build_style_profile,
    ensure_docx,
    parse_ksp,
    save_style_profile,
)

FIXTURES_DIR = Path(__file__).resolve().parent / "fixtures"
PROJECT_ROOT = Path(__file__).resolve().parent.parent
REAL_SCHEMA_PATH = PROJECT_ROOT / "storage" / "schema.sql"

THREE_FIXTURES = [
    "ksp_sample_1_single_table.docx",
    "ksp_sample_2_five_tables.docx",
    "ksp_sample_3_typos_case.docx",
]


# --- Б3.1: ensure_docx ---


def test_ensure_docx_returns_docx_path_unchanged():
    path = FIXTURES_DIR / "ksp_sample_1_single_table.docx"
    assert ensure_docx(path) == path


def test_ensure_docx_converts_real_doc_file():
    """На настоящем .doc (не .docx, переименованном в .doc) — реальная
    конвертация через LibreOffice, результат читаем python-docx."""
    doc_path = FIXTURES_DIR / "ksp_sample_1_single_table.doc"
    converted = ensure_docx(doc_path)

    assert converted.suffix == ".docx"
    assert converted.exists()

    document = Document(str(converted))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Краткосрочный" in text


def test_ensure_docx_rejects_unsupported_extension(tmp_path):
    fake = tmp_path / "not_a_ksp.txt"
    fake.write_text("просто текст")
    with pytest.raises(KSPConversionError):
        ensure_docx(fake)


def test_ensure_docx_fails_clearly_without_libreoffice(tmp_path, monkeypatch):
    import core.ksp_parser as ksp_parser_module

    monkeypatch.setattr(ksp_parser_module.shutil, "which", lambda name: None)
    fake_doc = tmp_path / "old.doc"
    fake_doc.write_bytes(b"content does not matter, soffice is never called")

    with pytest.raises(KSPConversionError, match="LibreOffice"):
        ensure_docx(fake_doc)


def test_ensure_docx_raises_on_timeout(tmp_path, monkeypatch):
    import subprocess

    import core.ksp_parser as ksp_parser_module

    def fake_run(cmd, **kwargs):
        raise subprocess.TimeoutExpired(cmd=cmd, timeout=60)

    monkeypatch.setattr(ksp_parser_module.subprocess, "run", fake_run)
    fake_doc = tmp_path / "slow.doc"
    fake_doc.write_bytes(b"...")

    with pytest.raises(KSPConversionError, match="60 секунд"):
        ensure_docx(fake_doc)


def test_ensure_docx_raises_on_nonzero_return_code(tmp_path, monkeypatch):
    import subprocess

    import core.ksp_parser as ksp_parser_module

    fake_result = subprocess.CompletedProcess(
        args=["soffice"], returncode=1, stdout=b"", stderr="конвертация не удалась".encode("utf-8")
    )
    monkeypatch.setattr(ksp_parser_module.subprocess, "run", lambda *a, **kw: fake_result)
    fake_doc = tmp_path / "broken.doc"
    fake_doc.write_bytes(b"...")

    with pytest.raises(KSPConversionError, match="конвертация не удалась"):
        ensure_docx(fake_doc)


def test_ensure_docx_raises_if_output_file_missing(tmp_path, monkeypatch):
    import subprocess

    import core.ksp_parser as ksp_parser_module

    fake_result = subprocess.CompletedProcess(args=["soffice"], returncode=0, stdout=b"", stderr=b"")
    monkeypatch.setattr(ksp_parser_module.subprocess, "run", lambda *a, **kw: fake_result)
    fake_doc = tmp_path / "vanishes.doc"
    fake_doc.write_bytes(b"...")

    with pytest.raises(KSPConversionError):
        ensure_docx(fake_doc)


# --- Б3.2: parse_ksp — три разных по вёрстке фикстуры не роняют парсер ---


@pytest.mark.parametrize("filename", THREE_FIXTURES)
def test_parse_ksp_returns_nonempty_structure(filename):
    result = parse_ksp(FIXTURES_DIR / filename)

    assert result["headings"], f"{filename}: заголовки не найдены"
    assert result["paragraphs"] or result["tables"], f"{filename}: ни абзацев, ни таблиц"
    assert result["tables"], f"{filename}: таблицы не найдены"


@pytest.mark.parametrize("filename", THREE_FIXTURES)
def test_parse_ksp_finds_lesson_plan_table_despite_different_layout(filename):
    """Главная проверка ловушки Б3.2: таблица "Ход урока" находится
    независимо от того, единственная это таблица в документе, одна из
    пяти, или её заголовки написаны капсом/с опечаткой/по-другому."""
    result = parse_ksp(FIXTURES_DIR / filename)
    lpt = result["lesson_plan_table"]

    assert lpt is not None, f"{filename}: таблица «Ход урока» не найдена"
    assert lpt["rows"], f"{filename}: у найденной таблицы нет строк с данными"

    expected_roles = {
        "etap_vremya",
        "deystviya_pedagoga",
        "deystviya_uchenika",
        "resursy",
        "ocenivanie",
    }
    assert expected_roles <= lpt["columns"].keys(), f"{filename}: {lpt['columns']}"


def test_parse_ksp_does_not_crash_on_document_without_matching_table(tmp_path):
    """Документ без узнаваемой таблицы «Ход урока» — штатная ситуация,
    не исключение: lesson_plan_table просто None."""
    document = Document()
    document.add_paragraph("Это не КСП, а случайный документ.")
    path = tmp_path / "random.docx"
    document.save(path)

    result = parse_ksp(path)

    assert result["lesson_plan_table"] is None
    assert result["tables"] == []
    assert result["paragraphs"] == ["Это не КСП, а случайный документ."]


def test_parse_ksp_raises_clear_error_on_corrupt_file(tmp_path):
    fake = tmp_path / "corrupt.docx"
    fake.write_bytes(b"this is not a real docx file, just garbage bytes")

    with pytest.raises(KSPParseError):
        parse_ksp(fake)


# --- проверка распознавания ролей колонок напрямую (опечатки, регистр) ---


@pytest.mark.parametrize(
    "cell_text,expected_role",
    [
        ("Этап урока / время", "etap_vremya"),
        ("ЭТАП УРОКА", "etap_vremya"),
        ("Действия педагога", "deystviya_pedagoga"),
        ("Дейтсвия  педагога", "deystviya_pedagoga"),  # опечатка + двойной пробел
        ("Действия ученика", "deystviya_uchenika"),
        ("Действия обучающегося", "deystviya_uchenika"),  # альтернативная формулировка
        ("Ресурсы", "resursy"),
        ("Ресурсы и материалы", "resursy"),
        ("Оценивание", "ocenivanie"),
        ("Оценка", "ocenivanie"),  # короткая форма, опечатка-подобная
    ],
)
def test_match_cell_role_recognizes_variants(cell_text, expected_role):
    assert _match_cell_role(cell_text) == expected_role


def test_match_cell_role_returns_none_for_unrelated_text():
    assert _match_cell_role("Механика, 10 класс") is None


# --- Б3.3: build_style_profile ---


class _FakeLLMClient:
    """Подменяет core.llm_client.LLMClient в тестах: не ходит в сеть,
    запоминает, что реально попало в промпт, и отдаёт заранее заданный
    JSON — как будто бы уже разобранный ответ LLM."""

    def __init__(self, response: dict):
        self.response = response
        self.calls: list[dict] = []

    async def complete_json(self, system, user, schema, max_retries=3):
        self.calls.append({"system": system, "user": user, "schema": schema})
        return self.response


async def test_build_style_profile_from_three_real_fixtures():
    parsed_list = [parse_ksp(FIXTURES_DIR / name) for name in THREE_FIXTURES]

    fake_llm = _FakeLLMClient(
        {
            "goal_phrasing": [
                "Все учащиеся смогут сформулировать закон сохранения импульса",
                "Все учащиеся смогут измерить влажность воздуха психрометром",
            ],
            "stage_structure": [
                {"stage": "Начало урока", "timing": "5 мин"},
                {"stage": "Основная часть", "timing": "25 мин"},
            ],
            "assessment_methods": ["Взаимооценивание", "Устный опрос"],
            "resources_used": ["Учебник", "Доска"],
        }
    )

    profile = await build_style_profile(parsed_list, llm_client=fake_llm)

    assert profile["goal_phrasing"]
    assert profile["stage_structure"]
    assert profile["assessment_methods"]
    assert profile["resources_used"]
    assert profile["raw_samples_count"] == 3

    # промпт реально построен из содержимого фикстур, а не из ничего
    assert len(fake_llm.calls) == 1
    sent_prompt = fake_llm.calls[0]["user"]
    assert "Документ 1" in sent_prompt
    assert "Документ 2" in sent_prompt
    assert "Документ 3" in sent_prompt


async def test_build_style_profile_empty_list_raises():
    with pytest.raises(ValueError):
        await build_style_profile([], llm_client=_FakeLLMClient({}))


async def test_build_style_profile_defaults_missing_fields_to_empty_list():
    parsed_list = [parse_ksp(FIXTURES_DIR / THREE_FIXTURES[0])]
    fake_llm = _FakeLLMClient({"goal_phrasing": ["одна цель"]})  # остальных полей нет

    profile = await build_style_profile(parsed_list, llm_client=fake_llm)

    assert profile["goal_phrasing"] == ["одна цель"]
    assert profile["stage_structure"] == []
    assert profile["assessment_methods"] == []
    assert profile["resources_used"] == []
    assert profile["raw_samples_count"] == 1


# --- Б3.4: save_style_profile ---


@pytest.fixture
def db_with_teacher(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute(
        "INSERT INTO teachers (id, name, subject) VALUES (1, 'Тестовый Учитель', 'физика')",
        db_path=db_path,
    )
    return db_path


def _sample_profile(raw_samples_count=3):
    return {
        "goal_phrasing": ["цель А"],
        "stage_structure": [{"stage": "начало", "timing": "5 мин"}],
        "assessment_methods": ["устный опрос"],
        "resources_used": ["учебник"],
        "raw_samples_count": raw_samples_count,
    }


def test_save_style_profile_inserts_new_row(db_with_teacher):
    save_style_profile(1, _sample_profile(), db_path=db_with_teacher)

    rows = query("SELECT * FROM style_profiles WHERE teacher_id = 1", db_path=db_with_teacher)
    assert len(rows) == 1
    assert json.loads(rows[0]["goal_phrasing"]) == ["цель А"]
    assert rows[0]["raw_samples_count"] == 3
    assert rows[0]["updated_at"] is not None


def test_save_style_profile_updates_not_duplicates(db_with_teacher):
    save_style_profile(1, _sample_profile(raw_samples_count=3), db_path=db_with_teacher)
    save_style_profile(1, _sample_profile(raw_samples_count=5), db_path=db_with_teacher)

    rows = query("SELECT * FROM style_profiles WHERE teacher_id = 1", db_path=db_with_teacher)
    assert len(rows) == 1, "повторная загрузка КСП тем же учителем не должна плодить дубли"
    assert rows[0]["raw_samples_count"] == 5


# =====================================================================
# Временная копия .doc -> .docx не остаётся на диске.
#
# ensure_docx кладёт результат конвертации в tempfile.mkdtemp и владельца
# ей не назначает. Пока её никто не убирал, каждый разобранный .doc
# оставлял на диске ПОЛНУЮ КОПИЮ КСП педагога — навсегда: путь к ней
# нигде не хранится, значит /delete_my_data о ней не знает.
# =====================================================================

import shutil
import tempfile as _tempfile


def _временные_папки_конвертации() -> set:
    from core.ksp_parser import _TEMP_DOCX_PREFIX

    корень = Path(_tempfile.gettempdir())
    return {p for p in корень.glob(f"{_TEMP_DOCX_PREFIX}*") if p.is_dir()}


def test_неудачная_конвертация_не_оставляет_временную_папку(tmp_path, monkeypatch):
    """Конвертация падает регулярно — сломанный профиль LibreOffice,
    таймаут, битый .doc. Каждая неудача оставляла пустую папку на диске,
    и на машине, которая работает круглосуточно, их накапливались сотни."""
    from core import ksp_parser

    было = _временные_папки_конвертации()
    monkeypatch.setattr(ksp_parser.shutil, "which", lambda _: "/usr/bin/soffice")

    class _Провал:
        returncode = 1
        stderr = b"profile is locked"

    monkeypatch.setattr(ksp_parser.subprocess, "run", lambda *a, **k: _Провал())

    исходник = tmp_path / "старый.doc"
    исходник.write_bytes(b"not really a doc")
    with pytest.raises(KSPConversionError):
        ksp_parser.ensure_docx(исходник)

    assert not (_временные_папки_конвертации() - было), (
        "после неудачной конвертации осталась временная папка"
    )


def test_разбор_doc_не_оставляет_копию_во_временной_папке(tmp_path, monkeypatch):
    """Удачная конвертация оставляла на диске ПОЛНУЮ КОПИЮ КСП педагога —
    навсегда: путь к ней нигде не хранится, значит /delete_my_data о ней
    не знает, и обещание удалить данные выполнялось не до конца.

    LibreOffice здесь не нужен: подменяем саму конвертацию, потому что
    проверяется уборка за ней, а не она сама (её проверяет
    test_ensure_docx_converts_real_doc_file)."""
    from core import ksp_parser

    было = _временные_папки_конвертации()

    настоящий_docx = (FIXTURES_DIR / "ksp_sample_1_single_table.docx").read_bytes()

    def _поддельная_конвертация(path):
        папка = Path(_tempfile.mkdtemp(prefix=ksp_parser._TEMP_DOCX_PREFIX))
        итог = папка / f"{Path(path).stem}.docx"
        итог.write_bytes(настоящий_docx)
        return итог

    monkeypatch.setattr(ksp_parser, "ensure_docx", _поддельная_конвертация)

    исходник = tmp_path / "старый.doc"
    исходник.write_bytes(b"not really a doc")
    результат = ksp_parser.parse_ksp(исходник)

    assert результат["tables"], "разбор ничего не прочитал — тест проверяет не то"
    assert not (_временные_папки_конвертации() - было), (
        "после разбора .doc осталась временная папка с копией документа педагога"
    )


def test_разбор_docx_не_трогает_папку_исходника(tmp_path):
    """Страховка от слишком жадной уборки: у .docx конвертации нет, и
    удалять родителя исходного файла нельзя ни при каких условиях —
    это storage/uploads со всеми файлами пользователя."""
    свой = tmp_path / "ksp.docx"
    свой.write_bytes((FIXTURES_DIR / "ksp_sample_1_single_table.docx").read_bytes())
    сосед = tmp_path / "не_трогать.txt"
    сосед.write_text("важное", encoding="utf-8")

    parse_ksp(свой)

    assert свой.exists(), "исходный .docx удалён — этого делать нельзя"
    assert сосед.exists(), "снесена папка исходника вместе с чужими файлами"
