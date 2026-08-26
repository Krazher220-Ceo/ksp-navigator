"""
tests/test_konspekt_builder.py — тесты core/konspekt_builder.py (блок К6).
"""

from datetime import date

from docx import Document

from core.konspekt_builder import (
    DRAFT_NOTICE_TEXT,
    TITLE_TEXT,
    build_konspekt_docx,
    build_konspekt_filename,
)

SAMPLE_CONTENT = {
    "tema": "Кинематика: путь и перемещение",
    "celi": ["Различать путь и перемещение"],
    "glavnoe": ["Путь — скаляр, перемещение — вектор"],
    "formuly": [{"formula": "s = v*t", "znachenie": "путь при равномерном движении"}],
    "primery": ["Решили задачу про поезд"],
    "terminy": [{"termin": "перемещение", "opredelenie": "вектор из начальной точки в конечную"}],
    "voprosy_dlya_samoproverki": ["Чем отличается путь от перемещения?"],
    "domashnee_zadanie": "§12, задачи 1-3",
}


def _all_text(document: Document) -> str:
    return "\n".join(p.text for p in document.paragraphs)


# --- честная пометка (К6, отдельная от F7 у КСП) ---


def test_notice_present_and_not_copied_from_ksp(tmp_path):
    out_path = tmp_path / "konspekt.docx"
    build_konspekt_docx(SAMPLE_CONTENT, out_path)

    document = Document(str(out_path))
    text = _all_text(document)
    assert DRAFT_NOTICE_TEXT in text
    # своя формулировка — не текст КСП дословно
    from core.docx_builder import DRAFT_NOTICE_TEXT as KSP_DRAFT_NOTICE_TEXT

    assert DRAFT_NOTICE_TEXT != KSP_DRAFT_NOTICE_TEXT
    assert "распознавания речи" in DRAFT_NOTICE_TEXT
    assert "утвержд" not in DRAFT_NOTICE_TEXT  # не "требует утверждения педагогом" — конспект никто не утверждает


def test_title_and_tema_present(tmp_path):
    out_path = tmp_path / "konspekt.docx"
    build_konspekt_docx(SAMPLE_CONTENT, out_path)

    document = Document(str(out_path))
    text = _all_text(document)
    assert TITLE_TEXT in text
    assert SAMPLE_CONTENT["tema"] in text


# --- разделы: только когда есть данные ---


def test_all_sections_with_data_are_rendered(tmp_path):
    out_path = tmp_path / "konspekt.docx"
    build_konspekt_docx(SAMPLE_CONTENT, out_path)

    document = Document(str(out_path))
    text = _all_text(document)
    assert "Различать путь и перемещение" in text
    assert "Путь — скаляр, перемещение — вектор" in text
    assert "s = v*t — путь при равномерном движении" in text
    assert "перемещение: вектор из начальной точки в конечную" in text
    assert "Решили задачу про поезд" in text
    assert "Чем отличается путь от перемещения?" in text
    assert "§12, задачи 1-3" in text


def test_empty_sections_are_not_rendered(tmp_path):
    """Б6.2/К4.1 честность распространяется и на .docx: пустое поле —
    просто нет раздела, не выдуманное содержимое и не пустой заголовок."""
    minimal_content = {
        "tema": "Тема без деталей",
        "celi": [],
        "glavnoe": ["Хоть что-то обязательное"],
        "formuly": [],
        "primery": [],
        "terminy": [],
        "voprosy_dlya_samoproverki": [],
        "domashnee_zadanie": "",
    }
    out_path = tmp_path / "konspekt_minimal.docx"
    build_konspekt_docx(minimal_content, out_path)

    document = Document(str(out_path))
    text = _all_text(document)
    assert "Цели:" not in text
    assert "Формулы:" not in text
    assert "Примеры:" not in text
    assert "Термины:" not in text
    assert "Вопросы для самопроверки:" not in text
    assert "Домашнее задание:" not in text
    assert "Главное:" in text


# --- имя файла ---


def test_build_konspekt_filename_transliterates_and_replaces_spaces():
    filename = build_konspekt_filename("Закон Ньютона", date(2025, 11, 12))
    assert filename == "Конспект_zakon_nyutona_2025-11-12.docx"


def test_build_konspekt_filename_strips_unsafe_characters():
    filename = build_konspekt_filename('Тема: "закон" / часть 2', date(2025, 1, 5))
    for bad_char in '"/:*?<>|\\':
        assert bad_char not in filename
    assert filename.endswith("_2025-01-05.docx")


def test_build_konspekt_filename_length_never_exceeds_100_chars():
    huge_topic = "Очень длинная тема урока " * 10
    filename = build_konspekt_filename(huge_topic, date(2025, 6, 1))
    assert len(filename) <= 100


def test_build_konspekt_filename_accepts_string_date():
    filename = build_konspekt_filename("Тема", "2025-03-14")
    assert filename == "Конспект_tema_2025-03-14.docx"


def test_topic_with_quotes_slashes_colons_does_not_break_actual_save(tmp_path):
    """Главная проверка (Б5.3-аналог для конспекта): реальное сохранение
    на диск с именем из темы с кавычками/слэшами/двоеточиями не падает."""
    content = dict(SAMPLE_CONTENT)
    content["tema"] = 'Тема: "закон" / часть 2'
    filename = build_konspekt_filename(content["tema"], date(2025, 1, 5))
    out_path = tmp_path / filename
    build_konspekt_docx(content, out_path)
    assert out_path.exists()
