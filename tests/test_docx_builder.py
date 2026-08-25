"""
tests/test_docx_builder.py — тесты core/docx_builder.py.

Использует реальные встроенные шаблоны из core/templates.py (блок Б4) —
не придуманные структуры, а те самые три, что попадут в бота.
"""

import json
import re
from datetime import date
from pathlib import Path

import pytest
from docx import Document

from core.db import execute, init_db
from core.docx_builder import (
    ADAPTACIYA_OOP_TEXT,
    DRAFT_NOTICE_TEXT,
    KRITERII_USPEHA_TITLE,
    MANDATORY_NOTICE_TEXT,
    MARGIN_CM,
    PAGE_HEIGHT_CM,
    PAGE_WIDTH_CM,
    PROVERENO_TEXT,
    RAZDATOCHNYE_MATERIALY_TITLE,
    _FIELD_LABELS,
    build_docx,
    build_filename,
)
from core.templates import get_template, load_builtin_templates

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REAL_SCHEMA_PATH = PROJECT_ROOT / "storage" / "schema.sql"

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


SAMPLE_CONTENT = {
    "organizaciya": "КГУ «Гимназия №27», Костанай",
    "razdel": "Механика",
    "fio_pedagoga": "Нургалиев М.К.",
    "data": "12.11.2025",
    "chasy": "1",
    "klass": "10А",
    "prisutstvuet": "24",
    "otsutstvuet": "2",
    "tema_uroka": 'Закон сохранения импульса: теория и практика "на пальцах"',
    "celi_obucheniya": "10.1.4.1 применять законы сохранения при решении задач",
    "celi_uroka": [
        "Все учащиеся смогут сформулировать закон сохранения импульса",
        "Большинство учащихся смогут решить задачу на неупругий удар",
    ],
    "hod_uroka": [
        {
            "etap_vremya": "Начало урока, 0-5 мин",
            "deystviya_pedagoga": "Приветствует, проверяет присутствующих",
            "deystviya_uchenika": "Готовятся к уроку",
            "resursy": "Классный журнал",
            "ocenivanie": "Устная похвала",
            "domashnee_zadanie": "§24, вопросы 1-3",
            "dop_literatura": "Перышкин, §18",
        },
        {
            "etap_vremya": "Середина урока, 5-30 мин",
            "deystviya_pedagoga": "Объясняет закон сохранения импульса",
            "deystviya_uchenika": "Слушают, решают задачи",
            "resursy": "Учебник, доска",
            "ocenivanie": "Взаимооценивание",
            "domashnee_zadanie": "",
            "dop_literatura": "",
        },
    ],
}


@pytest.fixture
def db_with_builtins(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute("INSERT INTO teachers (id, name, subject) VALUES (1, 'Т', 'физика')", db_path=db_path)
    load_builtin_templates(db_path=db_path)
    return db_path


def _template_ids(db_path) -> list[int]:
    from core.templates import list_templates

    return [t["id"] for t in list_templates(teacher_id=1, db_path=db_path)]


# --- Б5.1: build_docx на всех трёх встроенных шаблонах ---


@pytest.mark.parametrize("template_index", [0, 1, 2])
def test_build_docx_produces_reopenable_file_for_every_builtin_template(
    db_with_builtins, tmp_path, template_index
):
    template_id = _template_ids(db_with_builtins)[template_index]
    template = get_template(template_id, db_path=db_with_builtins)

    out_path = tmp_path / f"output_{template_index}.docx"
    result = build_docx(SAMPLE_CONTENT, template, out_path)

    assert result == out_path
    assert out_path.exists()

    # файл действительно валидный .docx, открывается обратно без исключений
    reopened = Document(str(out_path))
    assert len(reopened.tables) == 1
    assert len(reopened.paragraphs) > 0


def test_build_docx_accepts_structure_json_as_raw_string(db_with_builtins, tmp_path):
    """template["structure_json"] может прийти уже распарсенным (из
    core.templates) или сырой JSON-строкой (прямо из БД) — оба случая
    должны работать одинаково."""
    template_id = _template_ids(db_with_builtins)[0]
    template = get_template(template_id, db_path=db_with_builtins)

    import json

    raw_string_template = dict(template)
    raw_string_template["structure_json"] = json.dumps(template["structure_json"])

    out_path = tmp_path / "raw_string.docx"
    build_docx(SAMPLE_CONTENT, raw_string_template, out_path)
    assert out_path.exists()


# --- Б5.2: пометка "черновик" — F7, обязательна в каждом файле ---


@pytest.mark.parametrize("template_index", [0, 1, 2])
def test_draft_notice_present_as_first_paragraph(db_with_builtins, tmp_path, template_index):
    template_id = _template_ids(db_with_builtins)[template_index]
    template = get_template(template_id, db_path=db_with_builtins)

    out_path = tmp_path / f"draft_check_{template_index}.docx"
    build_docx(SAMPLE_CONTENT, template, out_path)

    document = Document(str(out_path))
    assert document.paragraphs[0].text == DRAFT_NOTICE_TEXT


def test_draft_notice_is_bold_and_red(db_with_builtins, tmp_path):
    template = get_template(_template_ids(db_with_builtins)[0], db_path=db_with_builtins)
    out_path = tmp_path / "draft_style.docx"
    build_docx(SAMPLE_CONTENT, template, out_path)

    document = Document(str(out_path))
    run = document.paragraphs[0].runs[0]
    assert run.bold is True
    assert run.font.color.rgb is not None


# --- Б5.1: оформление страницы ---


def test_page_setup_a4_margins_and_font(db_with_builtins, tmp_path):
    template = get_template(_template_ids(db_with_builtins)[0], db_path=db_with_builtins)
    out_path = tmp_path / "page_setup.docx"
    build_docx(SAMPLE_CONTENT, template, out_path)

    document = Document(str(out_path))
    section = document.sections[0]

    # A4 (21.0 x 29.7 см), допуск на округление EMU
    assert abs(section.page_width.cm - 21.0) < 0.05
    assert abs(section.page_height.cm - 29.7) < 0.05
    for margin in (section.left_margin, section.right_margin, section.top_margin, section.bottom_margin):
        assert abs(margin.cm - 2.0) < 0.05

    normal_font = document.styles["Normal"].font
    assert normal_font.name == "Times New Roman"
    assert normal_font.size.pt == 12


def test_table_has_visible_borders_style(db_with_builtins, tmp_path):
    template = get_template(_template_ids(db_with_builtins)[0], db_path=db_with_builtins)
    out_path = tmp_path / "borders.docx"
    build_docx(SAMPLE_CONTENT, template, out_path)

    document = Document(str(out_path))
    assert document.tables[0].style.name == "Table Grid"


def test_column_headers_are_bold(db_with_builtins, tmp_path):
    template = get_template(_template_ids(db_with_builtins)[0], db_path=db_with_builtins)
    out_path = tmp_path / "headers_bold.docx"
    build_docx(SAMPLE_CONTENT, template, out_path)

    document = Document(str(out_path))
    table = document.tables[0]
    # официальный шаблон: строка 8 - заголовки колонок (после 7 строк шапки+целей + "Ход урока")
    header_row = next(r for r in table.rows if r.cells[0].text.strip() == "Этап урока/ Время")
    for cell in header_row.cells:
        assert cell.paragraphs[0].runs[0].bold is True


# --- официальная форма: все обязательные поля приказа №130 присутствуют ---


def test_official_form_contains_all_required_fields(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    out_path = tmp_path / "official_full.docx"
    build_docx(SAMPLE_CONTENT, official, out_path)

    document = Document(str(out_path))
    full_text = "\n".join(p.text for p in document.paragraphs)
    full_text += "\n" + "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)

    # Формулировки — дословно по приложению 4 приказа МОН РК №130 в редакции
    # от 30.04.2025 № 98 (PLAN_STAGE1_EXT.md, блок Р1.2). Раньше здесь стояли
    # сокращения ("ФИО педагога:", "Кол-во присутствующих:") — они читались
    # нормально человеком, но приказ формулирует иначе, а это утверждённый
    # документ, не вольный пересказ.
    required_labels = [
        "Раздел:",
        "Фамилия, имя, отчество (при его наличии) педагога:",
        "Дата:",
        "Класс:",
        "Количество присутствующих:",
        "Количество отсутствующих:",
        "Тема урока:",
        "Цели обучения в соответствии с учебной программой:",
        "Цели урока:",
        "Ход урока",
        "Этап урока/ Время",
        "Действия педагога",
        "Действия ученика",
        "Ресурсы",
        "Оценивание",
    ]
    for label in required_labels:
        assert label in full_text, f"в официальной форме отсутствует обязательное поле: {label!r}"


def test_merged_rows_use_real_gridspan_not_duplicated_cells(db_with_builtins, tmp_path):
    """Полноширинные строки (Раздел, ФИО и т.п.) должны быть настоящим
    слиянием ячеек (w:gridSpan), а не пятью отдельными ячейками с
    одинаковым текстом — иначе в Word будут лишние границы посреди строки."""
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    out_path = tmp_path / "gridspan_check.docx"
    build_docx(SAMPLE_CONTENT, official, out_path)

    document = Document(str(out_path))
    table = document.tables[0]

    razdel_row = next(r for r in table.rows if r.cells[0].text.startswith("Раздел:"))
    tcs = razdel_row._tr.findall(f".//{W_NS}tc")
    assert len(tcs) == 1, "строка 'Раздел:' должна быть одной физически слитой ячейкой"

    klass_row = next(r for r in table.rows if r.cells[0].text.startswith("Класс:"))
    tcs = klass_row._tr.findall(f".//{W_NS}tc")
    assert len(tcs) == 3, "строка Класс/Присутств./Отсутств. должна давать ровно 3 физические ячейки"


def test_celi_uroka_list_renders_as_multiple_lines(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    out_path = tmp_path / "celi_list.docx"
    build_docx(SAMPLE_CONTENT, official, out_path)

    document = Document(str(out_path))
    table = document.tables[0]
    celi_uroka_row = next(r for r in table.rows if r.cells[0].text.startswith("Цели урока:"))
    cell_text = celi_uroka_row.cells[0].text
    for item in SAMPLE_CONTENT["celi_uroka"]:
        assert item in cell_text


def test_empty_hod_uroka_produces_one_placeholder_row_not_crash(db_with_builtins, tmp_path):
    content = dict(SAMPLE_CONTENT)
    content["hod_uroka"] = []
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    out_path = tmp_path / "empty_hod_uroka.docx"
    build_docx(content, official, out_path)  # не должно упасть

    document = Document(str(out_path))
    header_row_idx = next(
        i for i, r in enumerate(document.tables[0].rows) if r.cells[0].text.strip() == "Этап урока/ Время"
    )
    assert len(document.tables[0].rows) == header_row_idx + 2  # заголовок + 1 пустая строка-заглушка


def test_extended_template_has_seven_columns_in_hod_uroka(db_with_builtins, tmp_path):
    extended = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if "Развёрнутый" in t["name"]
    )
    out_path = tmp_path / "extended.docx"
    build_docx(SAMPLE_CONTENT, extended, out_path)

    document = Document(str(out_path))
    assert len(document.tables[0].columns) == 7


def test_gymnasium_template_omits_fields_it_does_not_declare(db_with_builtins, tmp_path):
    gymnasium = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if "Гимназия" in t["name"]
    )
    out_path = tmp_path / "gymnasium.docx"
    build_docx(SAMPLE_CONTENT, gymnasium, out_path)

    document = Document(str(out_path))
    full_text = "\n".join(cell.text for row in document.tables[0].rows for cell in row.cells)
    assert "ФИО педагога:" not in full_text
    assert "Часы:" in full_text  # известная метка, не "Chasy:"


# --- Б5.3: имя файла ---


def test_build_filename_transliterates_and_replaces_spaces():
    filename = build_filename("физика", "10А", "Закон Ньютона", date(2025, 11, 12))
    assert filename == "КСП_физика_10А_zakon_nyutona_2025-11-12.docx"


def test_build_filename_strips_unsafe_characters_from_topic():
    filename = build_filename("физика", "10А", 'Тема: "закон" / часть 2', date(2025, 1, 5))
    for bad_char in '"/:*?<>|\\':
        assert bad_char not in filename
    assert filename.endswith("_2025-01-05.docx")


def test_build_filename_length_never_exceeds_100_chars():
    huge_topic = "Очень длинная тема урока " * 10 + 'с кавычками "и слэшами/двоеточиями:тоже"'
    filename = build_filename("физика", "10А", huge_topic, date(2025, 6, 1))
    assert len(filename) <= 100


def test_build_filename_accepts_string_date():
    filename = build_filename("физика", "10А", "Тема", "2025-03-14")
    assert filename == "КСП_физика_10А_tema_2025-03-14.docx"


def test_topic_with_quotes_slashes_colons_does_not_break_actual_save(
    db_with_builtins, tmp_path
):
    """Главная проверка Б5.3: настоящее сохранение на диск с именем,
    построенным из темы с кавычками/слэшами/двоеточиями, не падает."""
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    tricky_topic = 'Закон "Ома": сила тока / напряжение * сопротивление?'
    content = dict(SAMPLE_CONTENT)
    content["tema_uroka"] = tricky_topic

    filename = build_filename("физика", "10А", tricky_topic, date(2025, 9, 1))
    out_path = tmp_path / filename

    result = build_docx(content, official, out_path)

    assert result.exists()
    # исходная (не транслитерированная) тема сохранилась в тексте документа —
    # ломается только имя файла, не содержимое
    document = Document(str(result))
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert tricky_topic in full_text


# --- Р1.1: порядок колонок "Ход урока" — дословно приложение 4 приказа
# №130 в редакции от 30.04.2025 № 98. До этого блока Ресурсы и Оценивание
# были перепутаны местами — расхождение с утверждённой формой. Читаем
# порядок из JSON-файлов НА ДИСКЕ, а не из константы в docx_builder.py:
# именно расхождение между константой и JSON-шаблонами и было причиной
# исходного бага (поправить надо было в двух местах, поправили в одном).

BUILTIN_TEMPLATES_DIR = PROJECT_ROOT / "storage" / "builtin_templates"

# Официальные 5 колонок обязаны идти именно в этом порядке в НАЧАЛЕ списка
# колонок любого встроенного шаблона (у "Развёрнутого образца" после них
# могут идти свои, неофициальные, домашнее задание/доп. литература — это
# нормально, они не часть приказа).
_OFFICIAL_COLUMN_ORDER = [
    "etap_vremya",
    "deystviya_pedagoga",
    "deystviya_uchenika",
    "ocenivanie",
    "resursy",
]


@pytest.mark.parametrize(
    "template_path", sorted(BUILTIN_TEMPLATES_DIR.glob("*.json")), ids=lambda p: p.name
)
def test_hod_uroka_column_order_matches_order_130_appendix_4(template_path):
    data = json.loads(template_path.read_text(encoding="utf-8"))
    hod_uroka_block = next(b for b in data["blocks"] if b["key"] == "hod_uroka")
    columns = hod_uroka_block["columns"]

    official_columns_present = [c for c in columns if c in _OFFICIAL_COLUMN_ORDER]
    assert official_columns_present == _OFFICIAL_COLUMN_ORDER, (
        f"{template_path.name}: порядок официальных колонок {official_columns_present} "
        f"не совпадает с приложением 4 приказа №130 {_OFFICIAL_COLUMN_ORDER}"
    )


# --- Р1.2: формулировки полей шапки — дословно приложение 4 приказа №130
# в редакции от 30.04.2025 № 98. Список-эталон выписан прямо здесь, а не
# импортирован из docx_builder.py: тест должен уметь поймать расхождение
# кода с приказом, а не сверять код сам с собой.

_ORDER_130_FIELD_LABELS = {
    "razdel": "Раздел:",
    "fio_pedagoga": "Фамилия, имя, отчество (при его наличии) педагога:",
    "data": "Дата:",
    "klass": "Класс:",
    "prisutstvuet": "Количество присутствующих:",
    "otsutstvuet": "Количество отсутствующих:",
    "tema_uroka": "Тема урока:",
    "celi_obucheniya": "Цели обучения в соответствии с учебной программой:",
    "celi_uroka": "Цели урока:",
}


def test_field_labels_match_order_130_appendix_4_verbatim():
    for field, official_label in _ORDER_130_FIELD_LABELS.items():
        assert _FIELD_LABELS[field] == official_label, (
            f"поле {field!r}: у нас {_FIELD_LABELS[field]!r}, "
            f"в приказе {official_label!r}"
        )


# --- Р1.3: строка "ПРОВЕРЕНО" — только у шаблонов, объявивших её ---


def test_provereno_line_appears_only_when_template_declares_it(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )

    # ни один встроенный шаблон пока не объявляет "provereno" — строки нет
    out_path = tmp_path / "no_provereno.docx"
    build_docx(SAMPLE_CONTENT, official, out_path)
    full_text = "\n".join(p.text for p in Document(str(out_path)).paragraphs)
    assert PROVERENO_TEXT not in full_text

    # шаблон, который её объявил, — строка появляется
    template_with_provereno = dict(official)
    structure = dict(template_with_provereno["structure_json"])
    structure["blocks"] = [
        dict(b, fields=[*b.get("fields", []), "provereno"]) if b["key"] == "shapka" else b
        for b in structure["blocks"]
    ]
    template_with_provereno["structure_json"] = structure

    out_path2 = tmp_path / "with_provereno.docx"
    build_docx(SAMPLE_CONTENT, template_with_provereno, out_path2)
    full_text2 = "\n".join(p.text for p in Document(str(out_path2)).paragraphs)
    assert PROVERENO_TEXT in full_text2

    # и не создаёт лишней строки "Provereno:" в самой таблице
    table = Document(str(out_path2)).tables[0]
    assert not any("provereno" in row.cells[0].text.lower() for row in table.rows)


# --- Р1.4: оба абзаца после "Ход урока" присутствуют дословно ---


def test_mandatory_notice_and_oop_note_both_present_verbatim(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    out_path = tmp_path / "notices.docx"
    build_docx(SAMPLE_CONTENT, official, out_path)

    full_text = "\n".join(p.text for p in Document(str(out_path)).paragraphs)
    assert MANDATORY_NOTICE_TEXT in full_text
    assert ADAPTACIYA_OOP_TEXT in full_text
    # порядок в документе — как в самом приказе: обязательность пунктов
    # плана раньше, примечание про ООП следом
    assert full_text.index(MANDATORY_NOTICE_TEXT) < full_text.index(ADAPTACIYA_OOP_TEXT)


# --- Р3.2: раздаточные материалы — необязательный раздел ---


def test_document_without_razdatochnye_materialy_looks_unchanged(db_with_builtins, tmp_path):
    """КГ Р3.2: content без карточек — документ выглядит ровно как раньше,
    без этого блока, а не с пустым заголовком раздела."""
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    out_path = tmp_path / "no_cards.docx"
    build_docx(SAMPLE_CONTENT, official, out_path)  # SAMPLE_CONTENT карточек не содержит

    full_text = "\n".join(p.text for p in Document(str(out_path)).paragraphs)
    assert RAZDATOCHNYE_MATERIALY_TITLE not in full_text
    assert KRITERII_USPEHA_TITLE not in full_text


def test_document_with_razdatochnye_materialy_renders_all_levels(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    content = dict(SAMPLE_CONTENT)
    content["razdatochnye_materialy"] = [
        {
            "uroven": "Базовый уровень",
            "metka": "зелёная метка",
            "zadanie": "Тележка массой 2 кг движется со скоростью 3 м/с...",
            "podskazka": "Используйте закон сохранения импульса: m1v1 = (m1+m2)v2",
            "reshenie": "v2 = 2 м/с",
        },
        {
            "uroven": "Продвинутый уровень",
            "metka": "синяя метка",
            "zadanie": "Создайте математическую модель движения двух шаров...",
            # решения нет намеренно — творческое задание, это не баг
        },
    ]
    content["kriterii_uspeha"] = [
        "Создана математическая модель для обоих типов ударов",
        "Проведён анализ распределения кинетической энергии",
    ]

    out_path = tmp_path / "with_cards.docx"
    build_docx(content, official, out_path)  # не должно упасть — открывается ниже

    full_text = "\n".join(p.text for p in Document(str(out_path)).paragraphs)
    assert RAZDATOCHNYE_MATERIALY_TITLE in full_text
    assert "Базовый уровень" in full_text
    assert "зелёная метка" in full_text
    assert "v2 = 2 м/с" in full_text
    assert "Продвинутый уровень" in full_text
    # у карточки без решения нет строки "Решение:" вообще — не выдумываем его
    advanced_idx = full_text.index("Продвинутый уровень")
    tail = full_text[advanced_idx:]
    next_card_or_end = tail.find("Критерии успеха")
    advanced_block = tail[: next_card_or_end if next_card_or_end != -1 else None]
    assert "Решение:" not in advanced_block
    assert KRITERII_USPEHA_TITLE in full_text
    assert "Создана математическая модель для обоих типов ударов" in full_text


# --- Р5.1: ценность для интеграции и предварительные знания ---


def test_cennost_and_predznaniya_appear_only_when_present(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )

    def _full_text(path):
        # "Ценность для интеграции" и "Предварительные знания" — строки
        # ТАБЛИЦЫ (Р5.1), не отдельные параграфы: document.paragraphs не
        # заходит внутрь ячеек таблицы, нужно явно добавлять их текст —
        # тот же паттерн, что в test_official_form_contains_all_required_fields.
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        return text

    without = build_docx(SAMPLE_CONTENT, official, tmp_path / "without.docx")
    full_text_without = _full_text(tmp_path / "without.docx")
    assert _FIELD_LABELS["cennost_integracii"] not in full_text_without
    assert _FIELD_LABELS["predvaritelnye_znaniya"] not in full_text_without

    content = dict(SAMPLE_CONTENT)
    content["cennost_integracii"] = "Созидание и новаторство"
    content["predvaritelnye_znaniya"] = "Основы кинематики"
    with_values = build_docx(content, official, tmp_path / "with.docx")
    full_text_with = _full_text(tmp_path / "with.docx")
    assert "Созидание и новаторство" in full_text_with
    assert "Основы кинематики" in full_text_with


# --- Р5.3: колонка "Дифференциация/ООП" и альбомная ориентация ---


def test_oop_column_appears_only_when_flag_set(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )

    without = build_docx(SAMPLE_CONTENT, official, tmp_path / "no_oop.docx")
    header_row_without = next(
        r for r in Document(str(without)).tables[0].rows if r.cells[0].text.startswith("Этап")
    )
    assert "Дифференциация/ООП" not in [c.text for c in header_row_without.cells]

    content = dict(SAMPLE_CONTENT)
    content["ima_oop"] = True
    content["hod_uroka"] = [
        {**row, "differenciaciya_oop": "Карточка с укрупнённым шрифтом"} for row in SAMPLE_CONTENT["hod_uroka"]
    ]
    with_oop = build_docx(content, official, tmp_path / "with_oop.docx")
    table = Document(str(with_oop)).tables[0]
    header_row = next(r for r in table.rows if r.cells[0].text.startswith("Этап"))
    header_texts = [c.text for c in header_row.cells]
    assert "Дифференциация/ООП" in header_texts
    assert len(header_texts) == 6  # было 5 официальных колонок, теперь 6

    full_text = "\n".join(c.text for row in table.rows for c in row.cells)
    assert "Карточка с укрупнённым шрифтом" in full_text


def test_six_columns_widths_sum_to_table_width_in_book_orientation(db_with_builtins, tmp_path):
    """Р5.3, явная ловушка из плана: колонка ООП увеличивает таблицу до
    шести колонок — ширины обязаны пересчитаться и уложиться в печатную
    область A4, не уехать за поля."""
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    content = dict(SAMPLE_CONTENT)
    content["ima_oop"] = True
    out_path = tmp_path / "oop_book.docx"
    build_docx(content, official, out_path)

    document = Document(str(out_path))
    table = document.tables[0]
    header_row = next(r for r in table.rows if r.cells[0].text.startswith("Этап"))
    tcs = header_row._tr.findall(f".//{W_NS}tc")
    assert len(tcs) == 6

    section = document.sections[0]
    expected_table_width_cm = (section.page_width.cm) - 2 * MARGIN_CM
    # w:tcW хранит ширину в "dxa" (твипы, 1/20 пункта), не в EMU — проверено
    # напрямую на реальном .docx перед тем, как доверять числу в тесте.
    # 1 см = 1440 твипов / 2.54 = 566.929... твипов.
    DXA_PER_CM = 1440 / 2.54
    total_width_cm = sum(
        int(tc.find(f"{W_NS}tcPr/{W_NS}tcW").get(f"{W_NS}w")) / DXA_PER_CM
        for tc in tcs
    )
    assert abs(total_width_cm - expected_table_width_cm) < 0.05  # округления, не расхождение по сути


def test_album_orientation_swaps_page_dimensions_and_widens_table(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    content = dict(SAMPLE_CONTENT)
    content["page_orientation"] = "album"
    out_path = tmp_path / "album.docx"
    build_docx(content, official, out_path)

    document = Document(str(out_path))
    section = document.sections[0]
    # альбомная — ширина страницы больше высоты (было наоборот в книжной)
    assert section.page_width.cm > section.page_height.cm
    assert round(section.page_width.cm, 1) == round(PAGE_HEIGHT_CM, 1)
    assert round(section.page_height.cm, 1) == round(PAGE_WIDTH_CM, 1)


def test_book_orientation_is_default_when_not_specified(db_with_builtins, tmp_path):
    official = next(
        t for t in [get_template(i, db_path=db_with_builtins) for i in _template_ids(db_with_builtins)]
        if t["is_official"] == 1
    )
    out_path = tmp_path / "default_orientation.docx"
    build_docx(SAMPLE_CONTENT, official, out_path)  # без page_orientation в content

    document = Document(str(out_path))
    section = document.sections[0]
    assert round(section.page_width.cm, 1) == round(PAGE_WIDTH_CM, 1)
    assert round(section.page_height.cm, 1) == round(PAGE_HEIGHT_CM, 1)
