"""
tests/test_ktp_parser.py — тесты core/ktp_parser.py.

Новый модуль этого блока (Б8) — не в исходном PLAN_STAGE1.md как
отдельный пункт, понадобился для команды /upload_ktp.
"""

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from core.db import execute, init_db, query
from core.ktp_parser import KTPParseError, parse_ktp_file, save_ktp_entries

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REAL_SCHEMA_PATH = PROJECT_ROOT / "storage" / "schema.sql"


def _make_xlsx(path: Path, rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)


def _make_docx_table(path: Path, rows: list[list]) -> None:
    doc = Document()
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = "Table Grid"
    for row_values in rows:
        row = table.add_row()
        for i, v in enumerate(row_values):
            row.cells[i].text = "" if v is None else str(v)
    doc.save(path)


HEADER = ["№", "Раздел", "Тема урока", "Цели обучения", "Часы", "Дата", "Четверть"]
DATA_ROWS = [
    [1, "Механика", "Кинематика точки", "10.1.1.1", 1, "01-05.09.25", 1],
    [2, "Механика", "Закон сохранения импульса", "10.1.4.1", 2, "06-10.11.25", 2],
    [3, "Механика", "Тема без кода", None, 1, "", 2],
]


# --- .xlsx ---


def test_parse_xlsx_extracts_all_rows(tmp_path):
    path = tmp_path / "ktp.xlsx"
    _make_xlsx(path, [HEADER] + DATA_ROWS)

    entries = parse_ktp_file(path)

    assert len(entries) == 3
    assert entries[0] == {
        "lesson_number": 1,
        "section": "Механика",
        "topic": "Кинематика точки",
        "objective_code": "10.1.1.1",
        "hours": 1,
        "planned_date": "01-05.09.25",
        "quarter": 1,
    }
    assert entries[2]["objective_code"] is None
    assert entries[2]["planned_date"] is None


def test_parse_xlsx_header_not_in_first_row(tmp_path):
    path = tmp_path / "ktp_offset.xlsx"
    _make_xlsx(path, [["КТП физика 10 класс"], [], HEADER] + DATA_ROWS)

    entries = parse_ktp_file(path)
    assert len(entries) == 3


# --- .docx ---


def test_parse_docx_extracts_all_rows(tmp_path):
    path = tmp_path / "ktp.docx"
    _make_docx_table(path, [HEADER] + DATA_ROWS)

    entries = parse_ktp_file(path)
    assert len(entries) == 3
    assert entries[1]["topic"] == "Закон сохранения импульса"
    assert entries[1]["objective_code"] == "10.1.4.1"


# --- ошибки ---


def test_unsupported_extension_raises():
    with pytest.raises(KTPParseError):
        parse_ktp_file("/tmp/whatever.pdf")


def test_corrupt_xlsx_raises_ktp_parse_error_not_raw_exception(tmp_path):
    """Раньше тут наружу утекал zipfile.BadZipFile вместо KTPParseError —
    в боте это ушло бы трейсбеком в чат (нарушение Б8.3)."""
    path = tmp_path / "corrupt.xlsx"
    path.write_bytes(b"this is not a real xlsx file")

    with pytest.raises(KTPParseError):
        parse_ktp_file(path)


def test_corrupt_docx_raises_ktp_parse_error(tmp_path):
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"this is not a real docx file")

    with pytest.raises(KTPParseError):
        parse_ktp_file(path)


def test_no_recognizable_header_raises(tmp_path):
    path = tmp_path / "no_header.xlsx"
    _make_xlsx(path, [["a", "b", "c"], [1, 2, 3]])

    with pytest.raises(KTPParseError):
        parse_ktp_file(path)


def test_empty_file_raises(tmp_path):
    path = tmp_path / "empty.xlsx"
    wb = Workbook()
    wb.save(path)

    with pytest.raises(KTPParseError):
        parse_ktp_file(path)


# --- save_ktp_entries: обнуление неизвестных кодов (внешний ключ) ---


@pytest.fixture
def db_with_teacher_and_one_code(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute("INSERT INTO teachers (id, name, subject) VALUES (1, 'Т', 'физика')", db_path=db_path)
    execute(
        "INSERT INTO curriculum_objectives (code, grade, section, subsection, description, thinking_level) "
        "VALUES ('10.1.4.1', 10, 'Механика', 'Законы сохранения', 'применять законы сохранения', 'применение')",
        db_path=db_path,
    )
    return db_path


def test_save_ktp_entries_nulls_unknown_objective_codes(db_with_teacher_and_one_code):
    entries = [
        {"lesson_number": 1, "section": "Механика", "topic": "A", "objective_code": "10.1.1.1", "hours": 1, "planned_date": None, "quarter": 1},
        {"lesson_number": 2, "section": "Механика", "topic": "B", "objective_code": "10.1.4.1", "hours": 1, "planned_date": None, "quarter": 1},
    ]
    result = save_ktp_entries(1, entries, db_path=db_with_teacher_and_one_code)

    # replaced=0 — у учителя не было прошлого КТП, заменять было нечего
    assert result == {"inserted": 2, "replaced": 0, "codes_not_found": 1}

    rows = query("SELECT topic, objective_code FROM ktp_entries WHERE teacher_id = 1 ORDER BY topic", db_path=db_with_teacher_and_one_code)
    by_topic = {r["topic"]: r["objective_code"] for r in rows}
    assert by_topic["A"] is None  # неизвестный код обнулён, вставка не упала по FK
    assert by_topic["B"] == "10.1.4.1"  # известный код сохранён


def test_save_ktp_entries_no_codes_at_all(db_with_teacher_and_one_code):
    entries = [{"lesson_number": 1, "section": "S", "topic": "T", "objective_code": None, "hours": 1, "planned_date": None, "quarter": 1}]
    result = save_ktp_entries(1, entries, db_path=db_with_teacher_and_one_code)
    assert result == {"inserted": 1, "replaced": 0, "codes_not_found": 0}


def test_save_ktp_entries_replaces_previous_ktp_instead_of_duplicating(
    db_with_teacher_and_one_code,
):
    """КТП один на учебный год: повторная загрузка — это исправленный
    файл, а не второй КТП. Раньше записи просто дописывались, и после
    второй загрузки в базе лежали две версии сразу."""
    db_path = db_with_teacher_and_one_code
    first = [{"lesson_number": 1, "section": "Механика", "topic": "Старая тема", "objective_code": None, "hours": 1, "planned_date": None, "quarter": 1}]
    second = [{"lesson_number": 1, "section": "Механика", "topic": "Исправленная тема", "objective_code": None, "hours": 1, "planned_date": None, "quarter": 1}]

    save_ktp_entries(1, first, db_path=db_path)
    result = save_ktp_entries(1, second, db_path=db_path)

    assert result["replaced"] == 1
    rows = query("SELECT topic FROM ktp_entries WHERE teacher_id = 1", db_path=db_path)
    assert [r["topic"] for r in rows] == ["Исправленная тема"]


def test_save_ktp_entries_does_not_touch_other_teachers_ktp(db_with_teacher_and_one_code):
    """Замена КТП ограничена своим учителем — чужие записи не трогаются."""
    db_path = db_with_teacher_and_one_code
    execute("INSERT INTO teachers (id, name, subject) VALUES (2, 'Другой', 'химия')", db_path=db_path)
    entry = [{"lesson_number": 1, "section": "S", "topic": "Чужая тема", "objective_code": None, "hours": 1, "planned_date": None, "quarter": 1}]
    save_ktp_entries(2, entry, db_path=db_path)

    save_ktp_entries(1, entry, db_path=db_path)
    save_ktp_entries(1, entry, db_path=db_path)

    rows = query("SELECT topic FROM ktp_entries WHERE teacher_id = 2", db_path=db_path)
    assert [r["topic"] for r in rows] == ["Чужая тема"]
