"""
tests/test_generator.py — тесты core/ksp_generator.py.

LLM мокается через _ScriptedLLMClient (свой, простой fake, а не
httpx.MockTransport — здесь важна логика generate_ksp, не транспорт;
транспорт уже покрыт tests/test_llm_client.py). Реальные вызовы к
провайдерам в тестах запрещены.

guess_objective_code проверяется на настоящем curriculum_seed.sql —
том самом, что попадёт в реальную базу через scripts/init_db.py.
"""

import json
from pathlib import Path

import pytest
from docx import Document

from core.db import execute, init_db, query
from core.ksp_generator import (
    MAX_VIDY_DEYATELNOSTI,
    KSPGenerationError,
    KSPValidationError,
    LessonOptions,
    _extract_minutes,
    _fill_header_fields,
    _validate_ksp_content,
    build_prompt,
    generate_and_save_ksp,
    generate_ksp,
    guess_objective_code,
    save_generated_ksp,
)
from core.templates import get_template, load_builtin_templates

PROJECT_ROOT = Path(__file__).resolve().parent.parent
REAL_SCHEMA_PATH = PROJECT_ROOT / "storage" / "schema.sql"
REAL_SEED_PATH = PROJECT_ROOT / "curriculum_seed.sql"

DURATION = 40

# Реплики и дескрипторы ниже намеренно "полновесные" — не для красоты, а
# потому что core/ksp_generator.py, блок Р2.3, теперь отклоняет короткие
# заглушки вместо реального содержания (deystviya_pedagoga короче 120
# символов, ocenivanie из одного слова). VALID_CONTENT — это "хороший
# ответ модели", который _validate_ksp_content обязан пропускать без
# повтора; раньше здесь стояли короткие названия действий ("Объясняет
# тему", "Устный опрос") — они и были тем самым образцом бедного
# содержания, из-за которого блок Р2 писался.
VALID_CONTENT = {
    "tema_uroka": "Закон сохранения импульса",
    "razdel": "Механика",
    "celi_obucheniya": "10.1.4.1 применять законы сохранения при решении задач",
    "celi_uroka": ["Все учащиеся смогут сформулировать закон сохранения импульса"],
    "hod_uroka": [
        {
            "etap": "Начало урока",
            "vremya": "0-5",
            "deystviya_pedagoga": (
                '"Добрый день! Сегодня разберём закон сохранения импульса. '
                'Представьте: вы стоите на льду с тяжёлым рюкзаком. Что '
                'произойдёт, если резко бросить рюкзак вперёд? Обсудите в '
                'парах и запишите предположение в тетрадь."'
            ),
            "deystviya_uchenika": "Слушают, обсуждают в парах, записывают предположение",
            "resursy": "Интерактивная доска",
            "ocenivanie": "Ученик формулирует предположение\nОбосновывает его физическим законом",
        },
        {
            "etap": "Основная часть",
            "vremya": "5-35",
            "deystviya_pedagoga": (
                '"Разберём формулу m1v1 = m2v2 на примере тележки и груза. '
                'Кто попробует объяснить, почему импульс системы сохраняется? '
                'Теперь решите задачу на карточке самостоятельно, потом '
                'сверим ответы вместе."'
            ),
            "deystviya_uchenika": "Решают задачи в парах, сверяют ответы",
            "resursy": "Карточки с задачами, калькулятор",
            "ocenivanie": "Верно применяет формулу сохранения импульса\nПолучает правильный числовой ответ",
        },
        {
            "etap": "Итог урока",
            "vremya": "35-40",
            "deystviya_pedagoga": (
                '"Подведём итог: что мы узнали о законе сохранения импульса? '
                'Поднимите зелёный кружок, если уверены в теме, жёлтый — '
                'если остались вопросы."'
            ),
            "deystviya_uchenika": "Формулируют вывод, проводят самооценку кружками",
            "resursy": "Цветные кружки",
            "ocenivanie": "Формулирует вывод своими словами\nЧестно оценивает своё понимание темы",
        },
    ],
}


class _ScriptedLLMClient:
    """Отдаёт заранее заданные ответы по очереди, запоминает все вызовы
    (system/user/schema) — для проверки, что промпт на повторе реально
    содержит причину провала, а не просто дублирует исходный запрос."""

    def __init__(self, responses: list[dict]):
        self.responses = list(responses)
        self.calls: list[dict] = []

    async def complete_json(self, system, user, schema, max_retries=3):
        self.calls.append({"system": system, "user": user, "schema": schema})
        if not self.responses:
            raise AssertionError("LLM вызван больше раз, чем задано сценарием теста")
        return self.responses.pop(0)


# --- Б6.1: build_prompt ---


def test_build_prompt_contains_all_required_pieces():
    prompt = build_prompt(
        topic="Закон сохранения импульса",
        razdel="Механика",
        objective_code="10.1.4.1",
        klass="10А",
        duration_minutes=DURATION,
    )
    assert "Закон сохранения импульса" in prompt
    assert "Механика" in prompt
    assert "10.1.4.1" in prompt
    assert "10А" in prompt
    assert f"{DURATION} мин" in prompt


def test_build_prompt_fetches_objective_description_from_db(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute(
        "INSERT INTO curriculum_objectives (code, grade, section, subsection, description, thinking_level) "
        "VALUES ('10.1.4.1', 10, 'Механика', 'Законы сохранения', "
        "'применять законы сохранения при решении задач', 'применение')",
        db_path=db_path,
    )

    prompt = build_prompt(
        "Тема", "Раздел", "10.1.4.1", "10А", DURATION, db_path=db_path
    )
    assert "применять законы сохранения при решении задач" in prompt


def test_build_prompt_without_objective_code_says_not_specified():
    prompt = build_prompt("Тема", "Раздел", None, "10А", DURATION)
    assert "не указан" in prompt


def test_build_prompt_includes_context_only_when_style_profile_present():
    without_profile = build_prompt("Тема", "Раздел", None, "10А", DURATION, style_profile=None)
    assert "КОНТЕКСТ" not in without_profile

    profile = {
        "goal_phrasing": ["Все учащиеся смогут..."],
        "stage_structure": [{"stage": "Начало", "timing": "5 мин"}],
        "assessment_methods": ["Устный опрос"],
        "resources_used": ["Учебник"],
    }
    with_profile = build_prompt("Тема", "Раздел", None, "10А", DURATION, style_profile=profile)
    assert "КОНТЕКСТ" in with_profile
    assert "Все учащиеся смогут..." in with_profile
    assert "Устный опрос" in with_profile


# =====================================================================
# М4.2 — компоновка промпта под префиксный кэш: стабильная часть строго
# впереди переменной (PLAN_STAGE2.md, блок М4.2)
# =====================================================================


def test_build_prompt_stable_prefix_precedes_variable_task_section():
    """Профиль стиля (стабилен в пределах учителя) и инструкция про
    раздатки (стабильна, зависит только от флага) обязаны идти РАНЬШЕ
    темы/раздела/класса/длительности (меняются на каждый вызов) — иначе
    любое изменение темы обнуляет кэш всего промпта целиком."""
    profile = {
        "goal_phrasing": ["Все учащиеся смогут..."],
        "stage_structure": [{"stage": "Начало", "timing": "5 мин"}],
        "assessment_methods": ["Устный опрос"],
        "resources_used": ["Учебник"],
    }
    prompt = build_prompt(
        topic="Закон Ома",
        razdel="Электричество",
        objective_code=None,
        klass="10А",
        duration_minutes=DURATION,
        style_profile=profile,
        include_razdatochnye_materialy=True,
    )
    context_pos = prompt.index("КОНТЕКСТ")
    razdatka_pos = prompt.index("разноуровневые")  # ключевое слово инструкции про раздатки
    task_pos = prompt.index("ЗАДАЧА:")
    topic_pos = prompt.index("Закон Ома")

    assert context_pos < razdatka_pos < task_pos < topic_pos, (
        "стабильная часть (контекст стиля, инструкция про раздатки) должна "
        "идти раньше переменной части (задача с темой урока)"
    )


def test_build_prompt_variable_tail_order_preserved():
    """Опции урока и текст учебника — оба переменные, оба обязаны идти
    ПОСЛЕ задачи (темы/раздела/класса), а не до неё."""
    prompt = build_prompt(
        topic="Закон Ома",
        razdel="Электричество",
        objective_code=None,
        klass="10А",
        duration_minutes=DURATION,
        textbook_text="Текст со страницы учебника про закон Ома.",
    )
    task_pos = prompt.index("ЗАДАЧА:")
    textbook_pos = prompt.index("Текст со страницы учебника")
    assert task_pos < textbook_pos


def test_build_prompt_reordering_did_not_drop_any_content():
    """Перестановка (М4.2) — только порядок, не содержание: с теми же
    аргументами, что и раньше, в промпте по-прежнему есть всё, что было
    (регрессия на существующий тест ЗАДАЧА, объективный код, контекст)."""
    profile = {
        "goal_phrasing": ["формулировка"],
        "stage_structure": [{"stage": "Этап", "timing": "10 мин"}],
        "assessment_methods": ["метод"],
        "resources_used": ["ресурс"],
    }
    prompt = build_prompt(
        topic="Тема",
        razdel="Раздел",
        objective_code="10.1.1.1",
        klass="10Б",
        duration_minutes=DURATION,
        style_profile=profile,
        include_razdatochnye_materialy=True,
        textbook_text="Текст учебника",
    )
    for expected in (
        "Тема", "Раздел", "10.1.1.1", "10Б", f"{DURATION} мин",
        "формулировка", "метод", "ресурс", "разноуровневые", "Текст учебника",
    ):
        assert expected in prompt, f"{expected!r} пропало из промпта после перестановки М4.2"


def test_build_prompt_empty_style_profile_is_treated_as_no_profile():
    """Свежесозданный учитель без загруженных КСП — профиль существует
    как объект, но все списки пустые. Это тоже штатный режим без
    стилизации, не КОНТЕКСТ с прочерками."""
    empty_profile = {
        "goal_phrasing": [],
        "stage_structure": [],
        "assessment_methods": [],
        "resources_used": [],
    }
    prompt = build_prompt("Тема", "Раздел", None, "10А", DURATION, style_profile=empty_profile)
    assert "КОНТЕКСТ" not in prompt


# --- Р6.2: текст со страницы учебника в промпте ---


def test_build_prompt_includes_textbook_text_when_given():
    prompt = build_prompt(
        "Тема", "Раздел", None, "10А", DURATION,
        textbook_text="Закон сохранения импульса гласит: суммарный импульс замкнутой системы...",
    )
    assert "Закон сохранения импульса гласит" in prompt


def test_build_prompt_without_textbook_text_unchanged():
    without = build_prompt("Тема", "Раздел", None, "10А", DURATION)
    with_empty = build_prompt("Тема", "Раздел", None, "10А", DURATION, textbook_text=None)
    with_blank = build_prompt("Тема", "Раздел", None, "10А", DURATION, textbook_text="   ")
    assert without == with_empty == with_blank


# --- Р5.1/Р5.2/Р5.3: LessonOptions ---


def test_lesson_options_truncates_vidy_deyatelnosti_to_max():
    """Р5.2, ловушка из плана: 'до трёх' — ограничение, не пожелание."""
    options = LessonOptions(vidy_deyatelnosti=["А", "Б", "В", "Г", "Д"])
    assert len(options.vidy_deyatelnosti) == MAX_VIDY_DEYATELNOSTI
    assert options.vidy_deyatelnosti == ["А", "Б", "В"]


def test_lesson_options_defaults_do_not_change_prompt():
    """Учитель, не тронувший расширенную форму, получает тот же промпт,
    что и до блока Р5 — LessonOptions() "по умолчанию" ничего не добавляет."""
    without = build_prompt("Тема", "Раздел", None, "10А", DURATION)
    with_default_options = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=LessonOptions())
    assert without == with_default_options


def test_prompt_includes_cennost_name_and_goal_when_chosen():
    options = LessonOptions(cennost_key="sozidaniye_novatorstvo")
    prompt = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=options)
    assert "Созидание и новаторство" in prompt
    assert "инновационное мышление" in prompt.lower()


def test_prompt_ignores_unknown_cennost_key():
    options = LessonOptions(cennost_key="несуществующий_ключ")
    prompt = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=options)
    without = build_prompt("Тема", "Раздел", None, "10А", DURATION)
    assert prompt == without  # неизвестный ключ тихо игнорируется, не ломает промпт


def test_prompt_includes_vidy_deyatelnosti():
    options = LessonOptions(vidy_deyatelnosti=["Групповая работа", "Финансовая грамотность"])
    prompt = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=options)
    assert "Групповая работа" in prompt
    assert "Финансовая грамотность" in prompt


def test_prompt_includes_oop_instruction_only_when_flag_set():
    with_oop = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=LessonOptions(ima_oop=True))
    without_oop = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=LessonOptions(ima_oop=False))
    assert "differenciaciya_oop" in with_oop
    assert "differenciaciya_oop" not in without_oop


def test_prompt_includes_sor_instruction_only_when_flag_set():
    prompt = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=LessonOptions(sor_instead_of_reflection=True))
    assert "СОР" in prompt


def test_prompt_includes_fizkultminutka_instruction_only_when_flag_set():
    prompt = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=LessonOptions(fizkultminutka=True))
    assert "физкультминутк" in prompt.lower()


def test_prompt_includes_tip_uroka_and_mezhpredmetnye_svyazi():
    options = LessonOptions(tip_uroka="Контроль", mezhpredmetnye_svyazi=["информатика", "математика"])
    prompt = build_prompt("Тема", "Раздел", None, "10А", DURATION, options=options)
    assert "Контроль" in prompt
    assert "информатика" in prompt
    assert "математика" in prompt


def test_fill_header_fields_stamps_cennost_predznaniya_orientation_oop():
    options = LessonOptions(
        cennost_key="edinstvo_solidarnost",
        predvaritelnye_znaniya="Основы кинематики",
        page_orientation="album",
        ima_oop=True,
    )
    filled = _fill_header_fields({}, teacher_id=1, klass="10А", generated_at="2026-09-01", options=options)
    assert filled["cennost_integracii"] == "Единство и солидарность"
    assert filled["predvaritelnye_znaniya"] == "Основы кинематики"
    assert filled["page_orientation"] == "album"
    assert filled["ima_oop"] is True


def test_fill_header_fields_without_options_does_not_add_new_keys():
    filled = _fill_header_fields({}, teacher_id=1, klass="10А", generated_at="2026-09-01")
    for key in ("cennost_integracii", "predvaritelnye_znaniya", "page_orientation", "ima_oop"):
        assert key not in filled


def test_fill_header_fields_stamps_organizaciya_from_teacher_school(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute(
        "INSERT INTO teachers (id, name, subject, school) VALUES (1, 'Т', 'физика', 'КГУ «Гимназия №27»')",
        db_path=db_path,
    )

    filled = _fill_header_fields({}, teacher_id=1, klass="10А", generated_at="2026-09-01", db_path=db_path)
    assert filled["organizaciya"] == "КГУ «Гимназия №27»"


def test_fill_header_fields_leaves_organizaciya_absent_when_school_not_set(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute("INSERT INTO teachers (id, name, subject) VALUES (1, 'Т', 'физика')", db_path=db_path)

    filled = _fill_header_fields({}, teacher_id=1, klass="10А", generated_at="2026-09-01", db_path=db_path)
    assert "organizaciya" not in filled


def test_fill_header_fields_does_not_overwrite_existing_organizaciya(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute(
        "INSERT INTO teachers (id, name, subject, school) VALUES (1, 'Т', 'физика', 'КГУ «Гимназия №27»')",
        db_path=db_path,
    )

    filled = _fill_header_fields(
        {"organizaciya": "Уже вписанное значение"},
        teacher_id=1,
        klass="10А",
        generated_at="2026-09-01",
        db_path=db_path,
    )
    assert filled["organizaciya"] == "Уже вписанное значение"


# --- Б6.2: валидация ---


def test_validate_accepts_correct_content():
    assert _validate_ksp_content(VALID_CONTENT, DURATION) == []


def test_validate_reports_missing_hod_uroka():
    content = {k: v for k, v in VALID_CONTENT.items() if k != "hod_uroka"}
    problems = _validate_ksp_content(content, DURATION)
    assert any("hod_uroka" in p for p in problems)


def test_validate_reports_empty_hod_uroka_list():
    content = dict(VALID_CONTENT)
    content["hod_uroka"] = []
    problems = _validate_ksp_content(content, DURATION)
    assert any("hod_uroka" in p for p in problems)


def test_validate_reports_empty_string_field():
    content = dict(VALID_CONTENT)
    content["tema_uroka"] = ""
    problems = _validate_ksp_content(content, DURATION)
    assert any("tema_uroka" in p for p in problems)


def test_validate_reports_empty_field_inside_hod_uroka_entry():
    content = json.loads(json.dumps(VALID_CONTENT))  # глубокая копия
    content["hod_uroka"][0]["deystviya_pedagoga"] = "   "
    problems = _validate_ksp_content(content, DURATION)
    assert any("deystviya_pedagoga" in p for p in problems)


def test_validate_reports_timing_far_off_duration():
    content = json.loads(json.dumps(VALID_CONTENT))
    content["hod_uroka"] = [
        {**content["hod_uroka"][0], "vremya": "0-5"}
    ]  # 5 минут вместо 40
    problems = _validate_ksp_content(content, DURATION)
    assert any("таймингов" in p for p in problems)


def test_validate_tolerates_5_minute_deviation():
    content = json.loads(json.dumps(VALID_CONTENT))
    content["hod_uroka"][-1]["vremya"] = "35-44"  # сумма 44 вместо 40, но в пределах ±5
    assert _validate_ksp_content(content, DURATION) == []


def test_validate_skips_timing_check_when_unparseable():
    content = json.loads(json.dumps(VALID_CONTENT))
    for stage in content["hod_uroka"]:
        stage["vremya"] = "в начале урока"  # без цифр вообще
    problems = _validate_ksp_content(content, DURATION)
    assert not any("таймингов" in p for p in problems)


def test_validate_does_not_mutate_or_patch_content():
    """Ловушка Б6.2: валидация только диагностирует, не подставляет
    заглушки вместо недостающих данных."""
    original = {k: v for k, v in VALID_CONTENT.items() if k != "hod_uroka"}
    before = json.loads(json.dumps(original))
    _validate_ksp_content(original, DURATION)
    assert original == before  # ничего не дописано и не изменено


def test_extract_minutes_range():
    assert _extract_minutes("0-5 мин") == 5
    assert _extract_minutes("5–35") == 30
    assert _extract_minutes("10 мин") == 10
    assert _extract_minutes("без цифр") is None
    assert _extract_minutes("") is None


# --- Б6.2: generate_ksp — сценарии из КГ ---


async def test_generate_ksp_valid_response_passes_without_repair():
    fake = _ScriptedLLMClient([VALID_CONTENT])
    result = await generate_ksp(
        teacher_id=1,
        topic="Закон сохранения импульса",
        razdel="Механика",
        objective_code=None,
        klass="10А",
        duration_minutes=DURATION,
        llm_client=fake,
    )
    assert result == VALID_CONTENT
    assert len(fake.calls) == 1


async def test_generate_ksp_missing_hod_uroka_triggers_one_repair_then_succeeds():
    broken = {k: v for k, v in VALID_CONTENT.items() if k != "hod_uroka"}
    fake = _ScriptedLLMClient([broken, VALID_CONTENT])

    result = await generate_ksp(
        teacher_id=1,
        topic="Тема",
        razdel="Раздел",
        objective_code=None,
        klass="10А",
        duration_minutes=DURATION,
        llm_client=fake,
    )

    assert result == VALID_CONTENT
    assert len(fake.calls) == 2
    assert "hod_uroka" in fake.calls[1]["user"]  # причина провала попала в повторный промпт


async def test_generate_ksp_placeholder_style_answer_triggers_one_repair_then_succeeds():
    """Р2.3, КГ: короткая реплика вместо прямой речи педагога и
    односложное 'Оценивание' — это заглушки, а не ошибка формата (все
    обязательные ключи на месте, просто содержание бедное). Модель
    получает ровно один шанс переписать это осмысленно."""
    placeholder = json.loads(json.dumps(VALID_CONTENT))  # глубокая копия
    placeholder["hod_uroka"][0]["deystviya_pedagoga"] = "Объясняет тему"  # < 120 симв.
    placeholder["hod_uroka"][1]["ocenivanie"] = "Опрос"  # одно слово

    fake = _ScriptedLLMClient([placeholder, VALID_CONTENT])

    result = await generate_ksp(
        teacher_id=1,
        topic="Тема",
        razdel="Раздел",
        objective_code=None,
        klass="10А",
        duration_minutes=DURATION,
        llm_client=fake,
    )

    assert result == VALID_CONTENT
    assert len(fake.calls) == 2
    repair_prompt = fake.calls[1]["user"]
    assert "deystviya_pedagoga" in repair_prompt
    assert "ocenivanie" in repair_prompt


async def test_generate_ksp_second_failure_raises_validation_error():
    broken = {k: v for k, v in VALID_CONTENT.items() if k != "hod_uroka"}
    fake = _ScriptedLLMClient([broken, broken])

    with pytest.raises(KSPValidationError):
        await generate_ksp(
            teacher_id=1,
            topic="Тема",
            razdel="Раздел",
            objective_code=None,
            klass="10А",
            duration_minutes=DURATION,
            llm_client=fake,
        )
    assert len(fake.calls) == 2  # ровно один повтор, не больше


async def test_generate_ksp_uses_teacher_style_profile_when_present(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute("INSERT INTO teachers (id, name, subject) VALUES (1, 'Т', 'физика')", db_path=db_path)
    execute(
        "INSERT INTO style_profiles (teacher_id, goal_phrasing, stage_structure, "
        "assessment_methods, resources_used, raw_samples_count) VALUES "
        "(1, ?, ?, ?, ?, 3)",
        (
            json.dumps(["Все учащиеся смогут применить закон"]),
            json.dumps([{"stage": "Начало", "timing": "5 мин"}]),
            json.dumps(["Взаимооценивание"]),
            json.dumps(["Учебник"]),
        ),
        db_path=db_path,
    )

    fake = _ScriptedLLMClient([VALID_CONTENT])
    await generate_ksp(
        teacher_id=1,
        topic="Тема",
        razdel="Раздел",
        objective_code=None,
        klass="10А",
        duration_minutes=DURATION,
        llm_client=fake,
        db_path=db_path,
    )

    assert "КОНТЕКСТ" in fake.calls[0]["user"]
    assert "Все учащиеся смогут применить закон" in fake.calls[0]["user"]


async def test_generate_ksp_without_style_profile_is_not_an_error(tmp_path):
    """Штатный режим: учитель без профиля стиля (ещё не загрузил КСП) —
    генерация всё равно проходит, просто без КОНТЕКСТ-раздела."""
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute("INSERT INTO teachers (id, name, subject) VALUES (1, 'Т', 'физика')", db_path=db_path)

    fake = _ScriptedLLMClient([VALID_CONTENT])
    result = await generate_ksp(
        teacher_id=1,
        topic="Тема",
        razdel="Раздел",
        objective_code=None,
        klass="10А",
        duration_minutes=DURATION,
        llm_client=fake,
        db_path=db_path,
    )
    assert result == VALID_CONTENT
    assert "КОНТЕКСТ" not in fake.calls[0]["user"]


# --- Б6.3: guess_objective_code на настоящем curriculum_seed.sql ---


@pytest.fixture
def db_with_real_seed(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute(
        "INSERT INTO teachers (id, name, subject) VALUES (1, 'Дмитрий Александрович', 'физика')",
        db_path=db_path,
    )
    from core.db import connect

    conn = connect(db_path)
    try:
        conn.executescript(REAL_SEED_PATH.read_text(encoding="utf-8"))
        conn.commit()
    finally:
        conn.close()
    return db_path


def test_guess_objective_code_exact_topic_from_real_seed(db_with_real_seed):
    code = guess_objective_code(teacher_id=1, topic="Закон сохранения импульса", db_path=db_with_real_seed)
    assert code == "10.1.4.1"


def test_guess_objective_code_is_case_insensitive_for_cyrillic(db_with_real_seed):
    """Ловушка: SQL LIKE в SQLite регистронезависим только для ASCII.
    guess_objective_code сравнивает в Python именно поэтому."""
    code = guess_objective_code(teacher_id=1, topic="закон сохранения импульса", db_path=db_with_real_seed)
    assert code == "10.1.4.1"


def test_guess_objective_code_made_up_topic_returns_none(db_with_real_seed):
    code = guess_objective_code(
        teacher_id=1, topic="Квантовая хромодинамика в 10 классе", db_path=db_with_real_seed
    )
    assert code is None


def test_guess_objective_code_empty_topic_returns_none(db_with_real_seed):
    assert guess_objective_code(teacher_id=1, topic="   ", db_path=db_with_real_seed) is None


def test_guess_objective_code_unknown_teacher_returns_none(db_with_real_seed):
    code = guess_objective_code(teacher_id=999, topic="Закон сохранения импульса", db_path=db_with_real_seed)
    assert code is None


# --- Б6.4: save_generated_ksp / generate_and_save_ksp ---


@pytest.fixture
def db_with_official_template(tmp_path):
    db_path = tmp_path / "test.db"
    init_db(db_path=db_path, schema_path=REAL_SCHEMA_PATH)
    execute("INSERT INTO teachers (id, name, subject) VALUES (1, 'Т', 'физика')", db_path=db_path)
    load_builtin_templates(db_path=db_path)
    official_id = query(
        "SELECT id FROM templates WHERE is_official = 1", db_path=db_path
    )[0]["id"]
    return db_path, official_id


def test_save_generated_ksp_writes_file_and_matching_db_row(db_with_official_template, tmp_path):
    db_path, template_id = db_with_official_template
    template = get_template(template_id, db_path=db_path)
    output_dir = tmp_path / "generated"

    result = save_generated_ksp(
        teacher_id=1,
        template_id=template_id,
        content=VALID_CONTENT,
        template=template,
        subject="физика",
        klass="10А",
        db_path=db_path,
        output_dir=output_dir,
    )

    docx_path = Path(result["docx_path"])
    assert docx_path.exists()
    assert docx_path.parent == output_dir

    row = query("SELECT * FROM generated_ksp WHERE id = ?", (result["id"],), db_path=db_path)[0]
    assert row["docx_path"] == str(docx_path)
    assert row["teacher_id"] == 1
    assert row["template_id"] == template_id

    # Ответ модели сохраняется полностью и без искажений...
    saved_content = json.loads(row["content_json"])
    for key, value in VALID_CONTENT.items():
        assert saved_content[key] == value
    # ...а сверх него дописываются поля шапки формы №130, которых в
    # ответе модели нет и быть не может (ФИО педагога, дата, класс).
    # Раньше здесь стояло строгое равенство с VALID_CONTENT — оно и
    # закрепляло баг: в документ уходил только ответ модели, и шапка
    # приказа №130 оставалась пустой.
    assert saved_content["fio_pedagoga"] == "Т"  # teachers.name из фикстуры
    assert saved_content["klass"] == "10А"
    assert saved_content["data"]

    # файл реально валиден
    document = Document(str(docx_path))
    assert len(document.tables) == 1


async def test_generated_document_has_no_empty_mandatory_header_cells(
    db_with_official_template, tmp_path
):
    """Стык Б6 -> Б5: то, что реально доезжает до документа.

    Раньше такого теста не было, и это позволило багу дожить до приёмки:
    test_official_form_contains_all_required_fields проверяет ЛЕЙБЛЫ и
    кормит docx_builder рукописным SAMPLE_CONTENT, где заполнено всё.
    Но генератор возвращает только 5 полей (схема ответа модели), а
    шапка формы №130 требует больше — и ФИО педагога, дата и класс
    выходили пустыми при каждой генерации. Здесь проверяются ЗНАЧЕНИЯ,
    и вход берётся ровно тот, что даёт настоящий конвейер."""
    db_path, template_id = db_with_official_template
    fake = _ScriptedLLMClient([VALID_CONTENT])

    result = await generate_and_save_ksp(
        teacher_id=1,
        template_id=template_id,
        topic="Закон сохранения импульса",
        razdel="Механика",
        subject="физика",
        klass="10А",
        duration_minutes=DURATION,
        llm_client=fake,
        db_path=db_path,
        output_dir=tmp_path / "generated",
    )

    table = Document(result["docx_path"]).tables[0]
    filled = {}
    for row in table.rows:
        text = row.cells[0].text
        if ":" in text:
            label, _, value = text.partition(":")
            filled.setdefault(label.strip(), value.strip())

    # Обязательные поля приказа №130, которые система знает сама и
    # обязана подставить: пустыми они остаться не могут. Лейбл — дословно
    # по приложению 4 приказа в редакции от 30.04.2025 № 98 (не "ФИО
    # педагога", PLAN_STAGE1_EXT.md, блок Р1.2).
    assert filled["Фамилия, имя, отчество (при его наличии) педагога"] == "Т"  # teachers.name
    assert filled["Класс"] == "10А"
    assert filled["Дата"]
    # ...и поля из ответа модели — на своих местах.
    assert filled["Тема урока"] == VALID_CONTENT["tema_uroka"]
    assert filled["Раздел"] == VALID_CONTENT["razdel"]

    # Количество присутствующих/отсутствующих ПУСТЫЕ намеренно: их вписывает
    # учитель на уроке, выдумывать их система не должна.
    klass_row = next(r for r in table.rows if r.cells[0].text.startswith("Класс:"))
    assert klass_row.cells[1].text.strip() == "Количество присутствующих:"


async def test_generate_and_save_ksp_full_pipeline(db_with_official_template, tmp_path):
    db_path, template_id = db_with_official_template
    fake = _ScriptedLLMClient([VALID_CONTENT])
    output_dir = tmp_path / "generated"

    result = await generate_and_save_ksp(
        teacher_id=1,
        template_id=template_id,
        topic="Закон сохранения импульса",
        razdel="Механика",
        subject="физика",
        klass="10А",
        duration_minutes=DURATION,
        llm_client=fake,
        db_path=db_path,
        output_dir=output_dir,
    )

    assert Path(result["docx_path"]).exists()
    rows = query("SELECT * FROM generated_ksp", db_path=db_path)
    assert len(rows) == 1


async def test_generate_and_save_ksp_unknown_template_raises(db_with_official_template):
    db_path, _template_id = db_with_official_template
    fake = _ScriptedLLMClient([VALID_CONTENT])

    with pytest.raises(KSPGenerationError):
        await generate_and_save_ksp(
            teacher_id=1,
            template_id=999999,
            topic="Тема",
            razdel="Раздел",
            subject="физика",
            klass="10А",
            duration_minutes=DURATION,
            llm_client=fake,
            db_path=db_path,
        )
